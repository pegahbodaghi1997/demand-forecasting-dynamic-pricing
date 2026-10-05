"""Archive Excel workbooks, validate rows, and load a weekly star schema."""

from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DATABASE = PROJECT_ROOT / "data" / "assessment.sqlite"
SCHEMA_PATH = PROJECT_ROOT / "sql" / "schema.sql"

COLUMNS = (
    "week",
    "product",
    "sales_units",
    "cost_of_goods_sold_cogs",
    "inventory_level",
    "demand_forecast",
    "price_per_unit",
    "gross_profit",
    "operating_profit",
    "revenue",
    "inventory_turnover_ratio",
    "gross_profit_margin",
    "revenue_growth",
)

ALIASES = {
    "week": ("week", "week_number", "week_id"),
    "product": ("product", "product_name", "commodity", "commodity_name"),
    "sales_units": ("sales_units", "units_sold", "sales_volume", "quantity_sold"),
    "cost_of_goods_sold_cogs": (
        "cost_of_goods_sold_cogs",
        "unit_cost",
        "cost_per_unit",
        "cogs_per_unit",
        "cogs",
    ),
    "inventory_level": ("inventory_level", "opening_inventory", "opening_stock", "stock"),
    "demand_forecast": ("demand_forecast", "historical_forecast", "forecast_units", "forecast"),
    "price_per_unit": ("price_per_unit", "unit_price", "selling_price"),
    "gross_profit": ("gross_profit",),
    "operating_profit": ("operating_profit",),
    "revenue": ("revenue", "sales_revenue"),
    "inventory_turnover_ratio": ("inventory_turnover_ratio", "inventory_turnover"),
    "gross_profit_margin": ("gross_profit_margin", "gross_margin"),
    "revenue_growth": ("revenue_growth", "revenue_growth_rate"),
}


def _header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


HEADER_LOOKUP = {
    _header_key(alias): canonical
    for canonical, aliases in ALIASES.items()
    for alias in aliases
}


def connect(database: Path = DEFAULT_DATABASE) -> sqlite3.Connection:
    database = Path(database)
    database.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA busy_timeout = 5000")
    connection.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))
    return connection


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _find_header(sheet: Any) -> tuple[int, dict[str, int], list[str]] | None:
    """Find a header near the top; tolerate column order, case, and extra columns."""
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True),
        start=1,
    ):
        mapping: dict[str, int] = {}
        names: list[str] = []
        ambiguous: set[str] = set()
        for index, value in enumerate(values):
            name = "" if value is None else str(value).strip()
            names.append(name or f"column_{index + 1}")
            canonical = HEADER_LOOKUP.get(_header_key(name))
            if canonical:
                if canonical in mapping:
                    ambiguous.add(canonical)
                mapping[canonical] = index
        if {"week", "product"}.issubset(mapping) and len(mapping) >= 4:
            if ambiguous:
                raise ValueError(
                    f"Ambiguous columns {sorted(ambiguous)!r} in sheet "
                    f"{sheet.title!r}, row {row_number}"
                )
            return row_number, mapping, names
    return None


def _stage_workbook(connection: sqlite3.Connection, run_id: int, workbook_bytes: bytes) -> int:
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    count = 0
    matched_sheets = 0
    try:
        with connection:
            for sheet in workbook.worksheets:
                header = _find_header(sheet)
                if header is None:
                    continue
                matched_sheets += 1
                header_row, mapping, names = header
                recognized_indices = set(mapping.values())
                for excel_row, values in enumerate(
                    sheet.iter_rows(min_row=header_row + 1, values_only=True),
                    start=header_row + 1,
                ):
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    raw = {
                        column: _cell_text(values[index]) if index < len(values) else None
                        for column, index in mapping.items()
                    }
                    extras = {
                        names[index] if index < len(names) else f"column_{index + 1}": _cell_text(value)
                        for index, value in enumerate(values)
                        if index not in recognized_indices and value is not None
                    }
                    connection.execute(
                        f"""
                        INSERT INTO raw_weekly (
                            run_id, sheet_name, excel_row, {", ".join(COLUMNS)}, extra_cells_json
                        ) VALUES ({", ".join("?" for _ in range(len(COLUMNS) + 4))})
                        """,
                        (
                            run_id,
                            sheet.title,
                            excel_row,
                            *(raw.get(column) for column in COLUMNS),
                            json.dumps(extras, ensure_ascii=False),
                        ),
                    )
                    count += 1
            if matched_sheets == 0:
                raise ValueError(
                    "No worksheet has Week and Product plus at least two recognized data columns"
                )
            connection.execute(
                "UPDATE ingestion_run SET raw_rows = ? WHERE run_id = ?", (count, run_id)
            )
    finally:
        workbook.close()
    return count


def _decimal(value: str | None) -> Decimal | None:
    if value is None:
        return None
    text = value.strip()
    if not text or text.lower() in {"n/a", "na", "null", "none", "-", "#n/a"}:
        return None
    text = text.replace("$", "").replace("€", "").replace("£", "").strip()
    if "," in text:
        if not re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?", text):
            raise ValueError(f"Ambiguous comma formatting: {value!r}")
        text = text.replace(",", "")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid numeric value: {value!r}") from exc
    if not number.is_finite():
        raise ValueError(f"Non-finite numeric value: {value!r}")
    return number


def _integer(value: str | None) -> int | None:
    number = _decimal(value)
    if number is None:
        return None
    if number != number.to_integral_value():
        raise ValueError(f"Expected a whole number: {value!r}")
    return int(number)


def _cents(value: str | None) -> int | None:
    number = _decimal(value)
    if number is None:
        return None
    return int((number * 100).to_integral_value(rounding=ROUND_HALF_UP))


def _issue(
    issues: list[tuple[int, int | None, str, str, str]],
    run_id: int,
    raw_id: int | None,
    severity: str,
    code: str,
    detail: str,
) -> None:
    issues.append((run_id, raw_id, severity, code, detail))


@dataclass(frozen=True)
class CleanRow:
    raw_id: int
    product_key: str
    product_name: str
    week: int
    sales: int
    cost_cents: int
    inventory: int | None
    forecast: int | None
    price_cents: int
    revenue_cents: int
    gross_cents: int
    operating_cents: int | None
    turnover: float | None
    sales_recovered: bool

    @property
    def key(self) -> tuple[str, int]:
        return self.product_key, self.week

    @property
    def signature(self) -> tuple[object, ...]:
        return (
            self.sales,
            self.cost_cents,
            self.inventory,
            self.forecast,
            self.price_cents,
            self.revenue_cents,
            self.gross_cents,
            self.operating_cents,
            self.turnover,
        )


def _clean_row(
    raw: sqlite3.Row,
    run_id: int,
    issues: list[tuple[int, int | None, str, str, str]],
) -> CleanRow | None:
    raw_id = raw["raw_id"]

    def warn(code: str, detail: str, severity: str = "warning") -> None:
        _issue(issues, run_id, raw_id, severity, code, detail)

    for column in (
        "cost_of_goods_sold_cogs",
        "price_per_unit",
        "revenue",
        "gross_profit",
        "operating_profit",
    ):
        try:
            amount = _decimal(raw[column])
        except ValueError:
            continue
        if amount is not None and amount * 100 != (amount * 100).to_integral_value():
            warn(
                "rounded_to_cents",
                f"{column} has more than two decimal places; rounded to cents",
            )

    try:
        week = _integer(raw["week"])
        if week is None or week < 1:
            raise ValueError("Week must be a positive integer")
        product_name = " ".join((raw["product"] or "").split())
        if not product_name:
            raise ValueError("Product is missing")
        cost = _cents(raw["cost_of_goods_sold_cogs"])
        price = _cents(raw["price_per_unit"])
        if cost is None or cost < 0:
            raise ValueError("Unit cost must be present and non-negative")
        if price is None or price <= 0:
            raise ValueError("Unit price must be present and positive")
        inventory = _integer(raw["inventory_level"])
        if inventory is not None and inventory < 0:
            raise ValueError("Opening inventory cannot be negative")
    except ValueError as exc:
        warn("rejected_required_field", str(exc), "error")
        return None

    try:
        sales = _integer(raw["sales_units"])
    except ValueError as exc:
        sales = None
        warn("invalid_sales", str(exc))
    recovered = sales is None
    if sales is None:
        try:
            source_revenue = _cents(raw["revenue"])
        except ValueError:
            source_revenue = None
        if source_revenue is None or source_revenue < 0:
            warn(
                "rejected_missing_sales",
                "Sales cannot be recovered without valid Revenue and Price_Per_Unit",
                "error",
            )
            return None
        implied = Decimal(source_revenue) / Decimal(price)
        nearest = int(implied.to_integral_value(rounding=ROUND_HALF_UP))
        if nearest < 0 or abs(nearest * price - source_revenue) > 1:
            warn(
                "rejected_missing_sales",
                "Revenue / unit price does not imply a whole number of units within one cent",
                "error",
            )
            return None
        sales = nearest
        warn("sales_recovered", f"Sales_Units recovered as {sales} from Revenue / Price")
    if sales < 0:
        warn("rejected_negative_sales", "Sales_Units cannot be negative", "error")
        return None

    try:
        forecast = _integer(raw["demand_forecast"])
        if forecast is not None and forecast < 0:
            raise ValueError("Demand_Forecast cannot be negative")
    except ValueError as exc:
        forecast = None
        warn("invalid_forecast", f"{exc}; stored as NULL")
    if forecast is None:
        warn("missing_forecast", "Historical baseline forecast is unavailable", "info")

    try:
        operating = _cents(raw["operating_profit"])
    except ValueError as exc:
        operating = None
        warn("invalid_operating_profit", f"{exc}; stored as NULL")

    try:
        turnover_number = _decimal(raw["inventory_turnover_ratio"])
        if turnover_number is not None and turnover_number < 0:
            raise ValueError("Inventory turnover cannot be negative")
        turnover = float(turnover_number) if turnover_number is not None else None
    except ValueError as exc:
        turnover = None
        warn("invalid_turnover", f"{exc}; stored as NULL")

    revenue = sales * price
    gross = sales * (price - cost)
    for column, calculated in (("revenue", revenue), ("gross_profit", gross)):
        try:
            source = _cents(raw[column])
        except ValueError as exc:
            source = None
            warn(f"invalid_{column}", str(exc))
        if source is not None and abs(source - calculated) > 1:
            warn(
                f"{column}_mismatch",
                f"Source {source} cents differs from calculated {calculated} cents; "
                "calculated value loaded",
            )

    try:
        source_margin = _decimal(raw["gross_profit_margin"])
        if source_margin is not None and revenue > 0:
            calculated_margin = Decimal(gross) / Decimal(revenue)
            if abs(source_margin - calculated_margin) > Decimal("0.0002"):
                warn("margin_mismatch", "Source margin differs from recalculated margin")
    except ValueError as exc:
        warn("invalid_margin", str(exc))

    if raw["revenue_growth"] is not None:
        try:
            _decimal(raw["revenue_growth"])
        except ValueError:
            warn("undefined_source_growth", "Source Revenue_Growth is non-finite or invalid", "info")

    if price <= cost:
        warn("price_at_or_below_cost", "Unit selling price is not above unit cost")
    if operating is not None and operating > gross:
        warn("operating_above_gross", "Operating profit exceeds calculated gross profit")
    if inventory is None:
        warn("missing_inventory", "Opening inventory is unavailable", "info")
    elif sales > inventory:
        warn(
            "sales_above_opening_inventory",
            "Sales exceed opening inventory; in-week receipts may explain this",
            "info",
        )
    if turnover == 0 and sales > 0:
        warn("zero_turnover_with_sales", "Source turnover is zero despite positive sales", "info")

    return CleanRow(
        raw_id=raw_id,
        product_key=product_name.casefold(),
        product_name=product_name,
        week=week,
        sales=sales,
        cost_cents=cost,
        inventory=inventory,
        forecast=forecast,
        price_cents=price,
        revenue_cents=revenue,
        gross_cents=gross,
        operating_cents=operating,
        turnover=turnover,
        sales_recovered=recovered,
    )


def _flag_trends(
    rows: list[CleanRow],
    run_id: int,
    issues: list[tuple[int, int | None, str, str, str]],
) -> None:
    previous: dict[str, CleanRow] = {}
    for row in sorted(rows, key=lambda item: (item.product_key, item.week)):
        prior = previous.get(row.product_key)
        previous[row.product_key] = row
        if prior is None or prior.week != row.week - 1:
            continue
        for label, old, new, minimum, factor in (
            ("unit_price", prior.price_cents, row.price_cents, 1, Decimal("1.20")),
            ("unit_cost", prior.cost_cents, row.cost_cents, 1, Decimal("1.20")),
            ("sales", prior.sales, row.sales, 100, Decimal("3")),
            ("inventory", prior.inventory, row.inventory, 100, Decimal("3")),
            ("historical_forecast", prior.forecast, row.forecast, 100, Decimal("2")),
        ):
            if old is None or new is None or old < minimum or new <= 0:
                continue
            ratio = Decimal(new) / Decimal(old)
            if ratio > factor or ratio < 1 / factor:
                _issue(
                    issues,
                    run_id,
                    row.raw_id,
                    "info",
                    f"trend_jump_{label}",
                    f"{label} changed from {old} to {new} between weeks "
                    f"{prior.week} and {row.week}; retained for review",
                )


def _load_clean(connection: sqlite3.Connection, run_id: int) -> None:
    raw_rows = connection.execute(
        "SELECT * FROM raw_weekly WHERE run_id = ? ORDER BY sheet_name, excel_row",
        (run_id,),
    ).fetchall()
    issues: list[tuple[int, int | None, str, str, str]] = []
    groups: dict[tuple[str, int], list[CleanRow]] = defaultdict(list)
    rejected = 0
    duplicates = 0
    for raw in raw_rows:
        clean = _clean_row(raw, run_id, issues)
        if clean is None:
            rejected += 1
        else:
            groups[clean.key].append(clean)

    accepted: list[CleanRow] = []
    for group in groups.values():
        first = group[0]
        if all(row.signature == first.signature for row in group[1:]):
            accepted.append(first)
            for duplicate in group[1:]:
                duplicates += 1
                _issue(
                    issues,
                    run_id,
                    duplicate.raw_id,
                    "info",
                    "duplicate_equivalent",
                    f"Same cleaned Product/Week values as raw row {first.raw_id}",
                )
        else:
            rejected += len(group)
            for conflicting in group:
                _issue(
                    issues,
                    run_id,
                    conflicting.raw_id,
                    "error",
                    "duplicate_conflict",
                    "Different values for the same Product and Week in this workbook",
                )

    if not accepted:
        raise ValueError("No valid, unambiguous weekly records could be loaded")
    _flag_trends(accepted, run_id, issues)

    with connection:
        connection.execute("DELETE FROM data_quality_issue WHERE run_id = ?", (run_id,))
        for row in accepted:
            connection.execute(
                """
                INSERT INTO dim_product (product_key, product_name)
                VALUES (?, ?)
                ON CONFLICT(product_key) DO NOTHING
                """,
                (row.product_key, row.product_name),
            )
            connection.execute(
                """
                INSERT INTO dim_week (week_id, relative_year, week_in_year)
                VALUES (?, ?, ?)
                ON CONFLICT(week_id) DO NOTHING
                """,
                (row.week, (row.week - 1) // 52 + 1, (row.week - 1) % 52 + 1),
            )
            product_id = connection.execute(
                "SELECT product_id FROM dim_product WHERE product_key = ?",
                (row.product_key,),
            ).fetchone()[0]
            connection.execute(
                """
                INSERT INTO fact_weekly_product (
                    product_id, week_id, sales_units, unit_cost_cents,
                    opening_inventory_units, historical_forecast_units,
                    unit_price_cents, revenue_cents, gross_profit_cents,
                    operating_profit_cents, turnover_ratio_source,
                    sales_recovered, source_run_id, source_raw_id
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(product_id, week_id) DO UPDATE SET
                    sales_units = excluded.sales_units,
                    unit_cost_cents = excluded.unit_cost_cents,
                    opening_inventory_units = excluded.opening_inventory_units,
                    historical_forecast_units = excluded.historical_forecast_units,
                    unit_price_cents = excluded.unit_price_cents,
                    revenue_cents = excluded.revenue_cents,
                    gross_profit_cents = excluded.gross_profit_cents,
                    operating_profit_cents = excluded.operating_profit_cents,
                    turnover_ratio_source = excluded.turnover_ratio_source,
                    sales_recovered = excluded.sales_recovered,
                    source_run_id = excluded.source_run_id,
                    source_raw_id = excluded.source_raw_id,
                    loaded_at = CURRENT_TIMESTAMP
                """,
                (
                    product_id,
                    row.week,
                    row.sales,
                    row.cost_cents,
                    row.inventory,
                    row.forecast,
                    row.price_cents,
                    row.revenue_cents,
                    row.gross_cents,
                    row.operating_cents,
                    row.turnover,
                    int(row.sales_recovered),
                    run_id,
                    row.raw_id,
                ),
            )
        connection.executemany(
            """
            INSERT INTO data_quality_issue
            (run_id, raw_id, severity, issue_code, detail)
            VALUES (?, ?, ?, ?, ?)
            """,
            issues,
        )
        connection.execute(
            """
            UPDATE ingestion_run
            SET status = 'loaded', loaded_at = CURRENT_TIMESTAMP, error_message = NULL,
                accepted_rows = ?, rejected_rows = ?, duplicate_rows = ?
            WHERE run_id = ?
            """,
            (len(accepted), rejected, duplicates, run_id),
        )


def load_excel(source: Path, database: Path = DEFAULT_DATABASE) -> dict[str, object]:
    source = Path(source).resolve()
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Only .xlsx and .xlsm workbooks are supported")
    workbook_bytes = source.read_bytes()
    digest = hashlib.sha256(workbook_bytes).hexdigest()
    connection = connect(database)
    try:
        run = connection.execute(
            "SELECT run_id, status, raw_rows FROM ingestion_run WHERE source_sha256 = ?",
            (digest,),
        ).fetchone()
        if run is None:
            with connection:
                cursor = connection.execute(
                    """
                    INSERT INTO ingestion_run
                    (source_sha256, source_name, source_workbook, status)
                    VALUES (?, ?, ?, 'staged')
                    """,
                    (digest, source.name, workbook_bytes),
                )
                run_id = cursor.lastrowid
            staged_rows = 0
        else:
            run_id = run["run_id"]
            staged_rows = run["raw_rows"]
            if run["status"] == "loaded":
                return run_summary(connection, run_id, skipped=True)
        try:
            if not staged_rows:
                _stage_workbook(connection, run_id, workbook_bytes)
            _load_clean(connection, run_id)
        except Exception as exc:
            with connection:
                connection.execute(
                    "UPDATE ingestion_run SET status = 'failed', error_message = ? "
                    "WHERE run_id = ?",
                    (str(exc), run_id),
                )
            raise
        return run_summary(connection, run_id, skipped=False)
    finally:
        connection.close()


def run_summary(
    connection: sqlite3.Connection, run_id: int, *, skipped: bool
) -> dict[str, object]:
    run = connection.execute(
        """
        SELECT run_id, source_name, source_sha256, status, raw_rows,
               accepted_rows, rejected_rows, duplicate_rows, error_message
        FROM ingestion_run WHERE run_id = ?
        """,
        (run_id,),
    ).fetchone()
    issue_counts = {
        row["issue_code"]: row["count"]
        for row in connection.execute(
            """
            SELECT issue_code, COUNT(*) AS count
            FROM data_quality_issue WHERE run_id = ?
            GROUP BY issue_code ORDER BY issue_code
            """,
            (run_id,),
        )
    }
    return {**dict(run), "skipped_existing": skipped, "issues": issue_counts}


def database_summary(database: Path = DEFAULT_DATABASE) -> dict[str, object]:
    connection = connect(database)
    try:
        counts = {
            table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            for table in (
                "ingestion_run",
                "raw_weekly",
                "dim_product",
                "dim_week",
                "fact_weekly_product",
                "data_quality_issue",
            )
        }
        runs = [
            dict(row)
            for row in connection.execute(
                """
                SELECT run_id, source_name, source_sha256, status, raw_rows,
                       accepted_rows, rejected_rows, duplicate_rows, error_message
                FROM ingestion_run ORDER BY run_id
                """
            )
        ]
        checks = {
            "foreign_key_violations": len(
                connection.execute("PRAGMA foreign_key_check").fetchall()
            ),
            "fact_without_raw": connection.execute(
                """
                SELECT COUNT(*) FROM fact_weekly_product AS f
                LEFT JOIN raw_weekly AS r ON r.raw_id = f.source_raw_id
                WHERE r.raw_id IS NULL
                """
            ).fetchone()[0],
            "revenue_formula_mismatches": connection.execute(
                """
                SELECT COUNT(*) FROM fact_weekly_product
                WHERE revenue_cents != sales_units * unit_price_cents
                """
            ).fetchone()[0],
            "gross_profit_formula_mismatches": connection.execute(
                """
                SELECT COUNT(*) FROM fact_weekly_product
                WHERE gross_profit_cents !=
                      sales_units * (unit_price_cents - unit_cost_cents)
                """
            ).fetchone()[0],
            "missing_historical_forecasts": connection.execute(
                """
                SELECT COUNT(*) FROM fact_weekly_product
                WHERE historical_forecast_units IS NULL
                """
            ).fetchone()[0],
            "recovered_sales_rows": connection.execute(
                """
                SELECT COUNT(*) FROM fact_weekly_product
                WHERE sales_recovered = 1
                """
            ).fetchone()[0],
        }
        return {"counts": counts, "runs": runs, "checks": checks}
    finally:
        connection.close()
